from __future__ import annotations

import os
import re
import unicodedata
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, Literal

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate, Spacer

from app.services.document_service import normalize_doc_type


ExportFormat = Literal["pdf", "docx"]


@dataclass(frozen=True)
class ExportedDocument:
    content: bytes
    media_type: str
    filename: str


@dataclass(frozen=True)
class DocumentSection:
    heading: str
    lines: tuple[str, ...]


DOCUMENT_TITLES = {
    "lean_canvas": "Lean Canvas",
    "bmc": "Business Model Canvas",
    "swot": "Phân tích SWOT",
    "product_plan": "Kế hoạch sản phẩm",
    "marketing_strategy": "Chiến lược tiếp thị",
    "funding_guide": "Kế hoạch gọi vốn",
}

FIELD_LABELS = {
    "problem": "Vấn đề", "solution": "Giải pháp",
    "unique_value_proposition": "Giá trị khác biệt", "unfair_advantage": "Lợi thế khó sao chép",
    "customer_segments": "Phân khúc khách hàng", "key_metrics": "Chỉ số chính",
    "channels": "Kênh tiếp cận", "cost_structure": "Cơ cấu chi phí",
    "revenue_streams": "Dòng doanh thu", "key_partners": "Đối tác chính",
    "key_activities": "Hoạt động chính", "key_resources": "Nguồn lực chính",
    "value_propositions": "Giá trị cung cấp", "customer_relationships": "Quan hệ khách hàng",
    "strengths": "Điểm mạnh", "weaknesses": "Điểm yếu", "opportunities": "Cơ hội",
    "threats": "Thách thức", "mvp_scope": "Phạm vi MVP", "features": "Tính năng ưu tiên",
    "timeline": "Lộ trình cột mốc", "target_audience": "Khách hàng mục tiêu",
    "key_messages": "Thông điệp chính", "budget_estimate": "Ngân sách dự kiến",
    "pitch_outline": "Dàn ý gọi vốn", "valuation_notes": "Ghi chú định giá",
    "funding_stage_recommendation": "Đề xuất giai đoạn gọi vốn",
}


class DocumentExportService:
    def export(
        self,
        *,
        startup_name: str,
        doc_type: str,
        content: dict[str, Any],
        file_format: ExportFormat,
    ) -> ExportedDocument:
        canonical_type = normalize_doc_type(doc_type)
        title = DOCUMENT_TITLES[canonical_type]
        sections = normalize_sections(content)
        basename = f"{slugify(startup_name or 'startup')}-{canonical_type.replace('_', '-')}"
        if file_format == "docx":
            return ExportedDocument(
                content=_build_docx(startup_name, title, sections),
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                filename=f"{basename}.docx",
            )
        if file_format == "pdf":
            return ExportedDocument(
                content=_build_pdf(startup_name, title, sections),
                media_type="application/pdf",
                filename=f"{basename}.pdf",
            )
        raise ValueError(f"Unsupported export format '{file_format}'.")

    def export_report(
        self,
        *,
        startup_name: str,
        sections: list[dict[str, Any]],
        file_format: ExportFormat,
    ) -> ExportedDocument:
        normalized = tuple(DocumentSection(section["title"], tuple(_report_lines(section.get("content", {})))) for section in sections)
        basename = f"{slugify(startup_name or 'startup')}-startup-report"
        if file_format == "docx":
            return ExportedDocument(_build_docx(startup_name, "Hồ sơ dự án", normalized), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", f"{basename}.docx")
        if file_format == "pdf":
            return ExportedDocument(_build_pdf(startup_name, "Hồ sơ dự án", normalized), "application/pdf", f"{basename}.pdf")
        raise ValueError(f"Unsupported export format '{file_format}'.")

    def export_pitch_deck(self, *, startup_name: str, slides: list[dict[str, str]]) -> ExportedDocument:
        return ExportedDocument(
            content=_build_pitch_pdf(startup_name, slides),
            media_type="application/pdf",
            filename=f"{slugify(startup_name or 'startup')}-pitch-deck.pdf",
        )


def normalize_sections(content: dict[str, Any]) -> tuple[DocumentSection, ...]:
    return tuple(
        DocumentSection(FIELD_LABELS.get(key, key.replace("_", " ").title()), tuple(_value_lines(value)))
        for key, value in content.items()
    )


def _value_lines(value: Any) -> list[str]:
    if value is None or value == "" or value == []:
        return ["Chưa có thông tin."]
    if isinstance(value, list):
        lines: list[str] = []
        for item in value:
            if isinstance(item, dict):
                lines.append(" · ".join(str(part) for part in item.values() if part not in (None, "")))
            else:
                lines.append(str(item))
        return [line for line in lines if line] or ["Chưa có thông tin."]
    if isinstance(value, dict):
        return [f"{FIELD_LABELS.get(key, key.replace('_', ' ').title())}: {item}" for key, item in value.items()]
    return [str(value)]


def _report_lines(content: dict[str, Any]) -> list[str]:
    lines: list[str] = []
    for key, value in content.items():
        if value in (None, "", [], {}):
            continue
        label = FIELD_LABELS.get(key, key.replace("_", " ").title())
        values = _value_lines(value)
        lines.extend(f"{label}: {item}" for item in values)
    return lines or ["Chưa có thông tin."]


def _build_docx(startup_name: str, title: str, sections: tuple[DocumentSection, ...]) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.4)
    heading = document.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(startup_name or "Startup chưa đặt tên")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(11)
    for item in sections:
        document.add_heading(item.heading, level=1)
        for line in item.lines:
            document.add_paragraph(line, style="List Bullet" if len(item.lines) > 1 else None)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _build_pdf(startup_name: str, title: str, sections: tuple[DocumentSection, ...]) -> bytes:
    font_name = _register_unicode_font()
    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=A4, rightMargin=22 * mm, leftMargin=22 * mm,
        topMargin=20 * mm, bottomMargin=20 * mm,
        title=title, author="AI Startup Coach",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("DocumentTitle", parent=styles["Title"], fontName=font_name, fontSize=20, leading=26, alignment=TA_CENTER)
    subtitle_style = ParagraphStyle("DocumentSubtitle", parent=styles["Normal"], fontName=font_name, fontSize=10, leading=14, alignment=TA_CENTER, textColor="#66746c")
    heading_style = ParagraphStyle("SectionHeading", parent=styles["Heading2"], fontName=font_name, fontSize=13, leading=18, spaceBefore=12, spaceAfter=6, textColor="#153f32")
    body_style = ParagraphStyle("Body", parent=styles["BodyText"], fontName=font_name, fontSize=10.5, leading=16)
    story: list[Any] = [Paragraph(_escape(title), title_style), Paragraph(_escape(startup_name or "Startup chưa đặt tên"), subtitle_style), Spacer(1, 8 * mm)]
    for item in sections:
        story.append(Paragraph(_escape(item.heading), heading_style))
        if len(item.lines) > 1:
            story.append(ListFlowable([ListItem(Paragraph(_escape(line), body_style)) for line in item.lines], bulletType="bullet", leftIndent=14))
        else:
            story.append(Paragraph(_escape(item.lines[0]), body_style))
    document.build(story)
    return output.getvalue()


def _build_pitch_pdf(startup_name: str, slides: list[dict[str, str]]) -> bytes:
    font_name = _register_unicode_font()
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=24 * mm,
        leftMargin=24 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title=f"Pitch Deck — {startup_name}",
        author="AI Startup Coach",
    )
    styles = getSampleStyleSheet()
    number_style = ParagraphStyle("SlideNumber", parent=styles["Normal"], fontName=font_name, fontSize=10, leading=14, textColor="#66746c")
    title_style = ParagraphStyle("SlideTitle", parent=styles["Title"], fontName=font_name, fontSize=30, leading=38, spaceAfter=14, textColor="#153f32")
    body_style = ParagraphStyle("SlideBody", parent=styles["BodyText"], fontName=font_name, fontSize=17, leading=26, textColor="#24332d")
    footer_style = ParagraphStyle("SlideFooter", parent=styles["Normal"], fontName=font_name, fontSize=10, leading=14, textColor="#66746c")
    story: list[Any] = []
    for index, slide in enumerate(slides):
        story.extend([
            Paragraph(f"{index + 1:02d}", number_style),
            Spacer(1, 12 * mm),
            Paragraph(_escape(str(slide.get("slide_title") or f"Slide {index + 1}")), title_style),
            Paragraph(_escape(str(slide.get("content") or "Chưa có nội dung.")), body_style),
            Spacer(1, 28 * mm),
            Paragraph(_escape(startup_name or "Startup chưa đặt tên"), footer_style),
        ])
        if index < len(slides) - 1:
            story.append(PageBreak())
    document.build(story)
    return output.getvalue()


def _register_unicode_font() -> str:
    font_name = "AIStartupCoachUnicode"
    if font_name in pdfmetrics.getRegisteredFontNames():
        return font_name
    candidates = [
        os.getenv("AI_COACH_FONT_PATH"),
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/segoeui.ttf",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            pdfmetrics.registerFont(TTFont(font_name, candidate))
            return font_name
    raise RuntimeError("No Unicode TrueType font was found for PDF export.")


def slugify(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii").lower()
    return re.sub(r"[^a-z0-9]+", "-", normalized).strip("-") or "startup"


def _escape(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
